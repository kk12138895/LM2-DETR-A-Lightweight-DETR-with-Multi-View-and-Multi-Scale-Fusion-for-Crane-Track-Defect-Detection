from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("output/documents/LCR-DETR_method_planning.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def keep_table_rows_together(table):
    for row_idx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        if row_idx == 0:
            tbl_header = tr_pr.find(qn("w:tblHeader"))
            if tbl_header is None:
                tr_pr.append(OxmlElement("w:tblHeader"))


def set_run_font(run, size=None, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, before=0, after=6, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p, after=6, line=1.1)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, 11, True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, 11)
    else:
        r = p.add_run(text)
        set_run_font(r, 11)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    for run in p.runs:
        set_run_font(run, 11)
    if not p.runs:
        r = p.add_run(text)
        set_run_font(r, 11)
    else:
        p.runs[0].text = text
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_paragraph(p, before=16 if level == 1 else 10, after=7 if level == 1 else 5, line=1.1)
    r = p.add_run(text)
    set_run_font(r, 16 if level == 1 else 13, True, "2E74B5" if level == 1 else "1F4D78")
    return p


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_table_widths(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    set_cell_margins(cell, 120, 160, 120, 160)
    p = cell.paragraphs[0]
    style_paragraph(p, after=3, line=1.1)
    r = p.add_run(title)
    set_run_font(r, 11, True, "1F4D78")
    p2 = cell.add_paragraph()
    style_paragraph(p2, after=0, line=1.1)
    r2 = p2.add_run(text)
    set_run_font(r2, 11)
    doc.add_paragraph()


def fill_table(table, rows, widths):
    table.style = "Table Grid"
    table.autofit = False
    set_table_widths(table, widths)
    for row_idx, row_data in enumerate(rows):
        row = table.rows[row_idx]
        for col_idx, text in enumerate(row_data):
            cell = row.cells[col_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, "F2F4F7")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(text)
            set_run_font(run, 10.5, row_idx == 0, "1F4D78" if row_idx == 0 else None)
    keep_table_rows_together(table)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = doc.add_paragraph()
    style_paragraph(title, after=4, line=1.1)
    title_run = title.add_run("LCR-DETR 论文主要研究工作与方法章节规划")
    set_run_font(title_run, 20, True, "0B2545")
    sub = doc.add_paragraph()
    style_paragraph(sub, after=12, line=1.1)
    sub_run = sub.add_run("用于落实老师要求：理清论文主要研究工作，并根据整理结果确定模块名称与章节名称。")
    set_run_font(sub_run, 11, False, "555555")

    add_callout(
        doc,
        "规划原则",
        "本文以 D-FINE 为基础，不在方法章节大段展开 FDR、GO-LSD、HGNetv2、NSA 或 C3K2 等前人方法。方法部分集中说明本文在轻量化特征编码、辅助视角融合、多尺度特征融合三个方向上的改动、动机和作用。"
    )

    add_heading(doc, "一、主要研究工作归纳", 1)
    add_heading(doc, "1. 轻量化特征编码", 2)
    add_body(doc, "目标：降低参数量和冗余计算，同时尽量保持 D-FINE 的检测性能。", "目标：")
    add_bullet(doc, "在 backbone 中使用 GhostConv-based G-CBA 替换 HGBlock 中的部分普通 CBA，减少卷积冗余。")
    add_bullet(doc, "在 encoder 中使用 MCSA 稀疏注意力，核心是将普通多头自注意力替换为基于 Native Sparse Attention 思想的稀疏注意力机制。")
    add_bullet(doc, "写作重点应放在“为什么轻量、为什么性能不明显下降”，不必展开前人模块的完整公式。")

    add_heading(doc, "2. 辅助视角跨视角门控融合", 2)
    add_body(doc, "目标：增强方向敏感缺陷、小目标、弱边界缺陷的表达能力。", "目标：")
    add_bullet(doc, "以后统一使用“辅助视角”表述，不在实验前把方法绑定为 180°。")
    add_bullet(doc, "模块名称统一为 Cross-view Gated Fusion Module (CGFM)，中文为“跨视角门控融合模块”。")
    add_bullet(doc, "差分特征 |x-y| 作为模块内部机制描述，用于引导通道门控和空间门控，但不放进模块名称。")

    add_heading(doc, "3. C3K2 多尺度特征融合", 2)
    add_body(doc, "目标：增强不同尺度特征之间的信息交互，提高复杂背景下的小缺陷和多尺度缺陷检测能力。", "目标：")
    add_bullet(doc, "D-FINE 的 Hybrid Encoder 已有 FPN/PAN 风格的多尺度融合路径。")
    add_bullet(doc, "本文不是发明 C3K2 或 FPN/PAN，而是将 C3K2 引入原有多尺度融合路径，替换部分融合 block。")
    add_bullet(doc, "建议命名为 C3K2-based Multi-scale Feature Fusion Module，比“残差可变卷积融合模块”更清楚。")

    add_heading(doc, "二、推荐模块命名", 1)
    rows = [
        ["研究主线", "建议英文名称", "建议中文名称", "说明"],
        ["轻量化", "Lightweight Feature Encoding Module (LFEM)", "轻量化特征编码模块", "统管 G-CBA 与 MCSA，符合老师“轻量化放一块”的组织要求。"],
        ["轻量化子模块", "GhostConv-based G-CBA", "基于 GhostConv 的轻量化卷积单元", "说明普通 CBA 被 G-CBA 替换，GhostConv 内部由主卷积与廉价操作生成特征。"],
        ["轻量化子模块", "MCSA Sparse Attention Encoder", "MCSA 稀疏注意力编码器", "核心是将普通多头自注意力替换为稀疏注意力，减少冗余 token 交互。"],
        ["辅助视角", "Cross-view Gated Fusion Module (CGFM)", "跨视角门控融合模块", "不再使用 Difference-guided；差分特征作为内部引导机制描述。"],
        ["多尺度", "C3K2-based Multi-scale Feature Fusion Module", "基于 C3K2 的多尺度特征融合模块", "强调引入 C3K2 改造 D-FINE 的多尺度融合路径。"],
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    fill_table(table, rows, [1.05, 2.05, 1.55, 1.85])

    add_heading(doc, "三、建议方法章节结构", 1)
    rows = [
        ["章节", "建议标题", "写作重点"],
        ["3.1", "Overall Architecture of LCR-DETR", "说明以 D-FINE 为基础，保留 FDR 和 GO-LSD 解码器；本文改动集中在轻量化、辅助视角、多尺度融合三部分。"],
        ["3.2", "Lightweight Feature Encoding Module", "把 G-CBA 和 MCSA 合并到轻量化主线下，避免方法看起来像四个松散模块。"],
        ["3.2.1", "GhostConv-based Lightweight HGBlock", "保留 G-CBA 图；弱化 HG-stage/HGBlock/ESE 原结构公式，重点解释替换普通卷积后的轻量化作用。"],
        ["3.2.2", "Sparse Attention Encoder with MCSA", "建议重画为普通注意力与稀疏注意力对比图；保留最关键公式，如 Top-k mask 后的 sparse attention。"],
        ["3.3", "Cross-view Gated Fusion Module", "统一写辅助视角，不写死 180°；说明差分特征、通道/空间门控和残差注入。"],
        ["3.4", "C3K2-based Multi-scale Feature Fusion Module", "说明在 D-FINE Hybrid Encoder 的 FPN/PAN 风格融合路径中引入 C3K2，增强多尺度特征交互。"],
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    fill_table(table, rows, [0.75, 2.35, 3.4])

    add_heading(doc, "四、原文处理建议", 1)
    add_bullet(doc, "D-FINE 的 FDR、GO-LSD 只在整体结构中一句说明，不在方法章节展开公式。")
    add_bullet(doc, "G-CBA 部分保留结构图，减少 HGBlock 原始流程公式的篇幅，把重点转到 GhostConv 如何降低冗余卷积。")
    add_bullet(doc, "MCSA 部分不要堆叠复杂的 Multi-Scale Compression 图和大量公式，改成说明“普通多头注意力到稀疏注意力”的替换逻辑。")
    add_bullet(doc, "CGFM 部分作为本文效果增强的关键模块保留完整描述，但模块名不再包含 Difference-guided，也不在模块名中绑定具体角度。")
    add_bullet(doc, "C3K2-Fusion 部分保留，但要明确 C3K2 来自已有结构，本文贡献是将其引入并适配 D-FINE 的多尺度融合路径。")

    add_heading(doc, "五、可直接给老师看的版本", 1)
    add_body(
        doc,
        "本文围绕 D-FINE 的轻量化部署和缺陷特征增强展开。方法部分归纳为三项研究工作：第一，轻量化特征编码，用 GhostConv-based G-CBA 减少卷积冗余，并用 MCSA 稀疏注意力降低冗余 token 计算；第二，辅助视角跨视角门控融合 CGFM，通过门控残差融合增强方向敏感和弱边界缺陷表达；第三，C3K2 多尺度特征融合，增强 D-FINE Hybrid Encoder 中不同尺度特征的信息交互。方法章节建议按“轻量化特征编码-跨视角门控融合-C3K2 多尺度融合”三部分组织，前人方法只作必要背景说明。"
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("LCR-DETR method planning")
    set_run_font(r, 9, False, "777777")

    doc.save(OUT)


if __name__ == "__main__":
    build()
